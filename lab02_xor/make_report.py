# -*- coding: utf-8 -*-
"""Формує звіт по ЛР-2 у .docx (стиль титулки — як у попередніх звітах студента)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT

DOC = Document()

# --- базовий стиль ---
st = DOC.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(14)
st.paragraph_format.line_spacing = 1.15
for s in DOC.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(2.5); s.right_margin = Cm(1.5)


def p(text="", *, align="left", bold=False, italic=False, size=14, space_after=6, font=None):
    par = DOC.add_paragraph()
    par.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT, "just": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    par.paragraph_format.space_after = Pt(space_after)
    if text:
        r = par.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if font:
            r.font.name = font
    return par


def heading(text):
    par = DOC.add_paragraph()
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(6)
    r = par.add_run(text); r.bold = True; r.font.size = Pt(14)
    return par


def code_block(lines):
    for ln in lines:
        par = DOC.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.left_indent = Cm(0.8)
        r = par.add_run(ln if ln else " ")
        r.font.name = "Courier New"; r.font.size = Pt(9)


def truth_table(rows, headers):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True; run.font.size = Pt(12)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(12)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    DOC.add_paragraph()
    return t


def caption(text):
    p(text, align="center", italic=True, size=12, space_after=12)


# ============================ ТИТУЛЬНА СТОРІНКА ============================
p("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align="center", space_after=0)
p("ЖИТОМИРСЬКА ПОЛІТЕХНІКА", align="center", bold=True, space_after=0)
p("Кафедра ММСА", align="center", space_after=0)
for _ in range(6):
    p(space_after=0)
p("ЛАБОРАТОРНА РОБОТА № 2", align="center", bold=True, size=16, space_after=0)
p("з дисципліни «Системи штучного інтелекту»", align="center", space_after=18)
p("Тема: НЕЙРОННА РЕАЛІЗАЦІЯ ЛОГІЧНИХ ФУНКЦІЙ AND, OR, XOR.",
  align="center", bold=True, space_after=0)
p("ПРОБЛЕМА XOR", align="center", bold=True, space_after=0)
for _ in range(7):
    p(space_after=0)
p("Виконав: студент групи ВТ-23-2", align="right", space_after=0)
p("Камінський Олексій Дмитрович", align="right", space_after=12)
p("Перевірив: Фант М.О.", align="right", space_after=0)
for _ in range(5):
    p(space_after=0)
p("Житомир – 2026", align="center", space_after=0)
DOC.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ============================ 1. МЕТА ============================
heading("1. Мета роботи")
p("Дослідити математичну модель нейрона; побудувати нейронну реалізацію логічних "
  "функцій AND та OR, з’ясувати суть проблеми XOR і реалізувати функцію XOR "
  "через композицію функцій OR та AND у вигляді двошарового персептрона.", align="just")

# ============================ 2. ТЕОРІЯ ============================
heading("2. Теоретичні відомості")
p("Формальний нейрон складається з суматора та функції активації. Суматор обчислює "
  "зважену суму входів разом із пороговим коефіцієнтом W₀:", align="just")
p("g(x) = W₁·x₁ + W₂·x₂ + W₀,", align="center", italic=True)
p("а порогова (одинична) функція активації переводить цю суму у бінарний вихід:", align="just")
p("f(g) = 1, якщо g ≥ 0;   f(g) = 0, якщо g < 0.", align="center", italic=True)
p("Рівняння g(x) = 0 задає у просторі входів розділяючу пряму: по один бік від неї нейрон "
  "видає 1, по інший — 0. Отже, одним нейроном можна реалізувати лише ті логічні функції, "
  "класи яких є лінійно роздільними. Для функцій AND та OR це виконується, а для XOR — ні: "
  "точки (0,1) і (1,0) належать до класу 1, а (0,0) і (1,1) — до класу 0, і жодна пряма не "
  "відокремлює одну пару від іншої. Це і є проблема XOR, через яку потрібен другий шар.", align="just")

# ============================ 3. ЗАВДАННЯ 1 ============================
heading("3. Завдання 1. Обчислювальний алгоритм функції xor(x₁, x₂) через or і and")
p("Нейрони OR та AND відрізняються лише пороговим коефіцієнтом W₀ (ваги входів однакові, "
  "W₁ = W₂ = 1):", align="just")
p("OR:  g = x₁ + x₂ − 0,5  →  розділяюча пряма x₁ + x₂ = 1/2;", align="center", italic=True)
p("AND: g = x₁ + x₂ − 1,5  →  розділяюча пряма x₁ + x₂ = 3/2.", align="center", italic=True)
p("Функція XOR будується як композиція вже наявних функцій: вона дорівнює одиниці тоді, "
  "коли хоча б один вхід одиничний (OR), але не обидва одночасно (NOT AND):", align="just")
p("xor(x₁, x₂) = and( or(x₁, x₂),  not and(x₁, x₂) ).", align="center", italic=True)
p("Реалізація мовою Python (файл xor_neuron.py):", align="just")
code_block([
    "def activation(g):",
    "    return 1 if g >= 0 else 0",
    "",
    "def neuron(x1, x2, w1, w2, w0):",
    "    return activation(w1 * x1 + w2 * x2 + w0)",
    "",
    "def or_gate(x1, x2):      # пряма x1 + x2 = 1/2",
    "    return neuron(x1, x2, w1=1.0, w2=1.0, w0=-0.5)",
    "",
    "def and_gate(x1, x2):     # пряма x1 + x2 = 3/2",
    "    return neuron(x1, x2, w1=1.0, w2=1.0, w0=-1.5)",
    "",
    "def xor_gate(x1, x2):",
    "    y1 = or_gate(x1, x2)         # перший шар: OR",
    "    y2 = and_gate(x1, x2)        # перший шар: AND",
    "    return neuron(y1, y2, w1=1.0, w2=-1.0, w0=-0.5)",
])
p()
p("Результати роботи програми наведено у таблицях 1–3.", align="just")

p("Таблиця 1 — Таблиця істинності функції OR", size=12, space_after=4)
truth_table([[0, 0, 0], [0, 1, 1], [1, 0, 1], [1, 1, 1]], ["x₁", "x₂", "OR"])

p("Таблиця 2 — Таблиця істинності функції AND", size=12, space_after=4)
truth_table([[0, 0, 0], [0, 1, 0], [1, 0, 0], [1, 1, 1]], ["x₁", "x₂", "AND"])

p("Таблиця 3 — Проміжні значення першого шару та результат XOR", size=12, space_after=4)
truth_table([[0, 0, 0, 0, 0], [0, 1, 1, 0, 1], [1, 0, 1, 0, 1], [1, 1, 1, 1, 0]],
            ["x₁", "x₂", "y₁ = OR", "y₂ = AND", "XOR"])

p("Отримана послідовність виходів 0, 1, 1, 0 повністю збігається з таблицею істинності "
  "функції XOR, отже алгоритм реалізовано правильно.", align="just")

DOC.add_picture("fig1_input_space.png", width=Cm(16.5))
DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Рисунок 1 — Розділяючі прямі для OR і AND та неможливість розділити XOR однією прямою")

# ============================ 4. ЗАВДАННЯ 2 ============================
heading("4. Завдання 2. Двошаровий персептрон і рівняння розділяючої прямої")
p("Структуру мережі наведено на рисунку 2. Перший (прихований) шар утворюють два нейрони "
  "з однаковими вагами входів W₁ = W₂ = 1 і різними порогами: нейрон y₁ реалізує OR "
  "(W₀ = −0,5), нейрон y₂ — AND (W₀ = −1,5). Другий шар має ваги W₁ = +1, W₂ = −1 і "
  "поріг W₀ = −0,5.", align="just")

DOC.add_picture("fig3_perceptron.png", width=Cm(15))
DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Рисунок 2 — Структура двошарового персептрона для функції XOR")

p("Рівняння розділяючих прямих першого шару (у просторі входів x₁, x₂):", align="just")
p("нейрон y₁ (OR):   x₁ + x₂ − 0,5 = 0,   тобто  x₁ + x₂ = 1/2;", align="center", italic=True)
p("нейрон y₂ (AND):  x₁ + x₂ − 1,5 = 0,   тобто  x₁ + x₂ = 3/2.", align="center", italic=True)
p("Ці дві прямі ділять площину на три смуги, і саме цим знімається проблема XOR: точки "
  "(0,1) та (1,0) потрапляють у середню смугу, а (0,0) і (1,1) — у крайні.", align="just")

p("Ключовий момент полягає в тому, що розділяюча пряма другого шару будується вже не в "
  "просторі входів, а в просторі виходів першого шару (y₁, y₂). Після перетворення першим "
  "шаром вихідні точки розміщуються так: (0,0) → (0,0); (0,1) та (1,0) → (1,0); "
  "(1,1) → (1,1). У цьому новому просторі клас «1» представлений єдиною точкою (1,0), "
  "і він уже лінійно роздільний. Нейрон другого шару обчислює", align="just")
p("g(y) = 1·y₁ + (−1)·y₂ − 0,5,", align="center", italic=True)
p("а шукане рівняння розділяючої прямої другого шару має вигляд:", align="just")
p("y₁ − y₂ − 0,5 = 0,   тобто   y₁ − y₂ = 1/2.", align="center", bold=True, italic=True)
p("Перевірка: для точки (1,0) маємо g = 1 − 0 − 0,5 = 0,5 ≥ 0 → вихід 1; для (0,0) "
  "g = −0,5 < 0 → вихід 0; для (1,1) g = 1 − 1 − 0,5 = −0,5 < 0 → вихід 0. "
  "Це і є таблиця істинності XOR.", align="just")

DOC.add_picture("fig2_hidden_space.png", width=Cm(11))
DOC.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Рисунок 3 — Простір виходів першого шару: класи стали лінійно роздільними")

# ============================ 5. ВИСНОВКИ ============================
heading("5. Висновки")
p("У ході лабораторної роботи досліджено математичну модель формального нейрона, що "
  "складається із суматора g(x) = ΣWᵢxᵢ + W₀ та порогової функції активації. Показано, що "
  "один нейрон задає у просторі ознак розділяючу пряму, а отже здатний реалізувати лише "
  "лінійно роздільні логічні функції. Для AND і OR достатньо одного нейрона: вони "
  "відрізняються тільки значенням порога W₀ (−1,5 та −0,5 відповідно) при однакових "
  "одиничних вагах входів.", align="just")
p("Функція XOR лінійно нероздільна, тому одним нейроном не реалізується — у цьому й полягає "
  "проблема XOR. Її розв’язано композицією: перший шар перетворює вхідні дані за допомогою "
  "нейронів OR та AND, після чого в новому просторі ознак (y₁, y₂) класи стають лінійно "
  "роздільними і задача розв’язується одним нейроном другого шару з рівнянням "
  "розділяючої прямої y₁ − y₂ = 1/2. Програмна реалізація підтвердила правильність "
  "побудови: отримані виходи 0, 1, 1, 0 повністю відповідають таблиці істинності XOR.", align="just")
p("Таким чином, додавання прихованого шару принципово розширює клас функцій, які здатна "
  "реалізувати нейронна мережа, — саме цим багатошарові мережі переважають одношаровий "
  "персептрон.", align="just")

OUT = "СШІ-ЛР-2-ВТ-23-2-Камінський.docx"
DOC.save(OUT)
print("Збережено", OUT)
