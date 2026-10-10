# -*- coding: utf-8 -*-
"""
Збирає звіт .docx (+ .pdf) з файлу report_content.md лабораторної роботи.

Використання:
    python build_report.py --dir <тека лаби> --num 5 \
        --theme "ДОСЛІДЖЕННЯ МЕТОДІВ АНСАМБЛЕВОГО НАВЧАННЯ"

Підтримуваний markdown: ## / ### заголовки, ```python / ```text блоки,
![підпис](файл.png), markdown-таблиці, абзаци, **жирний** інлайн.
"""
import argparse
import os
import re
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

STUDENT = "Камінський Олексій Дмитрович"
GROUP = "ВТ-23-2"
DISCIPLINE = "Системи штучного інтелекту"
PREFIX = "СШІ"
REPO = None
VARIANT = None
TEACHER = "____________________"
CITY_YEAR = "Житомир – 2026"

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "just": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def new_document():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.paragraph_format.line_spacing = 1.15
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
    return doc


def para(doc, text="", *, align="left", bold=False, italic=False, size=14, space_after=6):
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
    return p


def rich_para(doc, text, *, align="just", size=14, space_after=6):
    """Абзац із підтримкою **жирного** та `моноширинного` інлайну."""
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    p.paragraph_format.space_after = Pt(space_after)
    for chunk in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = p.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = p.add_run(chunk[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(11)
            continue
        else:
            run = p.add_run(chunk)
        run.font.size = Pt(size)
    return p


def title_page(doc, num, theme):
    para(doc, "МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align="center", space_after=0)
    para(doc, "ЖИТОМИРСЬКА ПОЛІТЕХНІКА", align="center", bold=True, space_after=0)
    para(doc, "Кафедра ММСА", align="center", space_after=0)
    for _ in range(6):
        para(doc, space_after=0)
    para(doc, f"ЛАБОРАТОРНА РОБОТА № {num}", align="center", bold=True, size=16, space_after=0)
    para(doc, f"з дисципліни «{DISCIPLINE}»", align="center", space_after=18)
    for line in theme.split("\n"):
        para(doc, line, align="center", bold=True, space_after=0)
    if VARIANT:
        # методичка вимагає зазначати варіант предметної області на титулці
        para(doc, space_after=0)
        for line in VARIANT.split("\n"):
            para(doc, line, align="center", italic=True, size=13, space_after=0)
    for _ in range(7):
        para(doc, space_after=0)
    para(doc, f"Виконав: студент групи {GROUP}", align="right", space_after=0)
    para(doc, STUDENT, align="right", space_after=12)
    para(doc, f"Перевірив: {TEACHER}", align="right", space_after=0)
    for _ in range(5):
        para(doc, space_after=0)
    para(doc, CITY_YEAR, align="center", space_after=0)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    if REPO:
        # методички ЛР4/5/7 прямо вимагають посилання на репозиторій у кожному звіті
        p_repo = doc.add_paragraph()
        p_repo.paragraph_format.space_after = Pt(10)
        r1 = p_repo.add_run("Репозиторій з кодом роботи: ")
        r1.bold = True
        r1.font.size = Pt(13)
        r2 = p_repo.add_run(REPO)
        r2.font.size = Pt(13)
        r2.font.name = "Courier New"


def add_code(doc, lines, mono_size=9):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Cm(0.8)
        run = p.add_run(line if line.strip() else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(mono_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, rows):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(cols):
            text = row[ci] if ci < len(row) else ""
            cells[ci].text = ""
            run = cells[ci].paragraphs[0].add_run(text.replace("**", ""))
            run.font.size = Pt(11 if cols <= 4 else 9)
            run.bold = ri == 0
            cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_picture(doc, path, caption, counter):
    # агент міг сам написати "Рисунок N — ..." у підписі; прибираємо, щоб не дублювалось
    caption = re.sub(r"^\s*рис(унок|\.)?\s*\d*\s*[.,—\-–:]*\s*", "", caption, flags=re.I).strip()
    if not os.path.exists(path):
        para(doc, f"[рисунок не знайдено: {os.path.basename(path)}]", align="center", italic=True)
        return counter
    doc.add_picture(path, width=Cm(15.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, f"Рисунок {counter} — {caption}", align="center", italic=True, size=12, space_after=12)
    return counter + 1


def render(doc, md_text, base_dir):
    lines = md_text.split("\n")
    i = 0
    fig = 1
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # горизонтальна лінія markdown (---, ***, ___) — візуальний роздільник, у .docx не потрібен
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            i += 1
            continue

        # заголовки
        if stripped.startswith("### "):
            para(doc, stripped[4:].strip(), bold=True, size=13, space_after=4)
            i += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[3:].strip())
            run.bold = True
            run.font.size = Pt(14)
            i += 1
            continue
        if stripped.startswith("# "):
            para(doc, stripped[2:].strip(), bold=True, size=15, space_after=8)
            i += 1
            continue

        # блок коду / виводу
        if stripped.startswith("```"):
            lang = stripped[3:].strip().lower()
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code(doc, block, mono_size=9 if lang == "python" else 8.5)
            continue

        # рисунок
        m = re.match(r"^!\[(.*)\]\((.+)\)\s*$", stripped)
        if m:
            fig = add_picture(doc, os.path.join(base_dir, m.group(2)), m.group(1), fig)
            i += 1
            continue

        # таблиця
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                is_sep = cells and all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c)
                if not is_sep:
                    width = len(rows[0]) if rows else len(cells)
                    if rows and (len(cells) < width or not cells[0]):
                        # перенесення рядка таблиці: доклеюємо текст до попереднього рядка
                        text = " ".join(c for c in cells if c)
                        if text:
                            idx = min(len(rows[-1]) - 1, max(0, len(cells) - 1))
                            rows[-1][idx] = (rows[-1][idx] + " " + text).strip()
                    else:
                        rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # списки
        if re.match(r"^[-*] ", stripped) or re.match(r"^\d+[.)] ", stripped):
            p = doc.add_paragraph(style="List Bullet" if stripped[0] in "-*" else "List Number")
            p.paragraph_format.space_after = Pt(2)
            content = re.sub(r"^([-*]|\d+[.)])\s+", "", stripped)
            for chunk in re.split(r"(\*\*[^*]+\*\*)", content):
                if not chunk:
                    continue
                run = p.add_run(chunk[2:-2] if chunk.startswith("**") else chunk)
                run.bold = chunk.startswith("**")
                run.font.size = Pt(13)
            i += 1
            continue

        # звичайний абзац (склеюємо до порожнього рядка)
        start = i
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,3} |```|!\[|\||[-*] |\d+[.)] )", lines[i].strip()
        ):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            rich_para(doc, " ".join(buf))
        if i == start:
            # рядок схожий на таблицю/картинку, але не розпізнаний жодною гілкою —
            # виводимо як звичайний текст і ОБОВʼЯЗКОВО рухаємось далі (інакше вічний цикл)
            rich_para(doc, stripped)
            i += 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--num", required=True)
    ap.add_argument("--theme", required=True)
    ap.add_argument("--src", default="report_content.md")
    ap.add_argument("--teacher", default=None)
    ap.add_argument("--discipline", default=None, help="назва дисципліни на титулці")
    ap.add_argument("--prefix", default=None, help="префікс імені файлу, напр. СШІ або МАПЗ")
    ap.add_argument("--repo", default=None, help="посилання на GitHub-репозиторій (вимога методичок)")
    ap.add_argument("--variant", default=None, help="варіант предметної області для титулки")
    ap.add_argument("--no-pdf", action="store_true")
    args = ap.parse_args()

    global TEACHER, DISCIPLINE, PREFIX, REPO, VARIANT
    if args.teacher:
        TEACHER = args.teacher
    if args.discipline:
        DISCIPLINE = args.discipline
    if args.prefix:
        PREFIX = args.prefix
    if args.repo:
        REPO = args.repo
    if args.variant:
        VARIANT = args.variant

    base = os.path.abspath(args.dir)
    md_path = os.path.join(base, args.src)
    if not os.path.exists(md_path):
        raise SystemExit(f"Немає {md_path}")

    doc = new_document()
    title_page(doc, args.num, args.theme)
    with open(md_path, encoding="utf-8") as f:
        render(doc, f.read(), base)

    out_name = f"{PREFIX}-ЛР-{args.num}-{GROUP}-Камінський.docx"
    out_path = os.path.join(base, out_name)
    doc.save(out_path)
    print("Збережено", out_path)

    if not args.no_pdf:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", out_path, "--outdir", base],
            check=False, capture_output=True, timeout=300,
        )
        pdf = out_path.replace(".docx", ".pdf")
        print("PDF:", pdf, "OK" if os.path.exists(pdf) else "НЕ СТВОРЕНО")


if __name__ == "__main__":
    main()
